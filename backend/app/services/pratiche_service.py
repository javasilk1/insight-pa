import json
import os
import re
from io import BytesIO
from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document
from docx.shared import Pt

from models.pratica import AnalisiPratica, DocumentoPratica, PraticaSintesi, RichiestaRelazione

ID_RE = re.compile(r"^P-\d{3}$")
FILE_RE = re.compile(r"^[\w\-]+\.pdf$")


def _default_dir() -> Path:
    env = os.getenv("DEMO_PRATICHE_DIR")
    if env:
        return Path(env)
    # docker: /demo_data montato dal compose; locale: cartella demo_data nella root del repo
    for candidate in (Path("/demo_data/pratiche"), Path(__file__).resolve().parents[3] / "demo_data" / "pratiche"):
        if candidate.is_dir():
            return candidate
    return Path("/demo_data/pratiche")


class PraticheService:
    """Legge le pratiche demo da disco. Sarà sostituito dal repository su DB + output dell'agente."""

    def __init__(self, base_dir: Path | None = None):
        self.base = base_dir or _default_dir()

    def _folder(self, pratica_id: str) -> Path | None:
        if not ID_RE.match(pratica_id):
            return None
        folder = self.base / pratica_id
        return folder if (folder / "ground_truth.json").is_file() else None

    def _gt(self, folder: Path) -> dict:
        return json.loads((folder / "ground_truth.json").read_text(encoding="utf-8"))

    def lista(self) -> list[PraticaSintesi]:
        out = []
        for folder in sorted(self.base.glob("P-*")):
            if not (folder / "ground_truth.json").is_file():
                continue
            gt = self._gt(folder)
            out.append(PraticaSintesi(
                id=gt["id"], indirizzo=gt["immobile"]["indirizzo"], intestatario=gt["immobile"]["intestatario"],
                scenario=gt["scenario"], esito=gt["esito_atteso"], n_difformita=len(gt["difformita_attese"]),
                n_documenti=len(gt["documenti"]),
            ))
        return out

    def analisi(self, pratica_id: str) -> AnalisiPratica | None:
        folder = self._folder(pratica_id)
        if not folder:
            return None
        gt = self._gt(folder)
        documenti = []
        for f in gt["documenti"]:
            tipo = re.sub(r"^\d+_", "", f.removesuffix(".pdf")).replace("_", " ")
            documenti.append(DocumentoPratica(
                file=f, tipo=tipo, pagine=len(PdfReader(folder / f).pages),
                url=f"/api/pratiche/{pratica_id}/documenti/{f}",
            ))
        return AnalisiPratica(
            id=gt["id"], origine="demo_ground_truth", esito=gt["esito_atteso"], sintesi=gt["sintesi"],
            immobile=gt["immobile"], documenti=documenti, dati_estratti=gt.get("dati_estratti_attesi", {}),
            difformita=gt["difformita_attese"], verifiche=gt.get("verifiche_attese", []),
        )

    def documento(self, pratica_id: str, file: str) -> Path | None:
        folder = self._folder(pratica_id)
        if not folder or not FILE_RE.match(file):
            return None
        path = folder / file
        return path if path.is_file() else None

    def relazione_docx(self, analisi: AnalisiPratica, req: RichiestaRelazione) -> bytes:
        confermate = [d for d in analisi.difformita if d.id in set(req.confermate)]
        im = analisi.immobile
        doc = Document()
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_heading("Relazione di verifica dello stato legittimo", 0)
        doc.add_paragraph().add_run(
            "BOZZA generata da InsightPA - da verificare, integrare e sottoscrivere a cura del tecnico."
        ).italic = True
        if analisi.origine == "demo_ground_truth":
            doc.add_paragraph("Pratica dimostrativa con dati fittizi.")

        doc.add_heading("1. Immobile", 1)
        catasto = f"Foglio {im.get('foglio')}, Particella {im.get('particella')}" + (f", Sub. {im['sub']}" if im.get("sub") else "")
        for label, value in [("Indirizzo", im.get("indirizzo")), ("Piano", im.get("piano")),
                             ("Intestatario", im.get("intestatario")), ("Dati catastali", catasto)]:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(str(value))

        doc.add_heading("2. Documenti esaminati", 1)
        for d in analisi.documenti:
            doc.add_paragraph(f"{d.tipo.capitalize()} ({d.file}, {d.pagine} pag.)", style="List Bullet")

        doc.add_heading("3. Esito della verifica", 1)
        if confermate:
            doc.add_paragraph(f"Sono state riscontrate {len(confermate)} difformità, descritte di seguito.")
        else:
            doc.add_paragraph("Non sono state riscontrate difformità tra lo stato di fatto e lo stato legittimo.")
        for v in analisi.verifiche:
            doc.add_paragraph(f"{v.descrizione} {v.valutazione} ({v.riferimento_normativo}).")

        if confermate:
            doc.add_heading("4. Difformità riscontrate", 1)
            for d in confermate:
                doc.add_heading(f"{d.id} - {d.tipo.replace('_', ' ').capitalize()} (ambito {d.ambito.replace('_', ' ')}, gravità {d.gravita})", 2)
                doc.add_paragraph(d.descrizione)
                if d.entita:
                    doc.add_paragraph(f"Entità: {d.entita}")
                doc.add_paragraph(f"Valutazione: {d.valutazione}")
                doc.add_paragraph(f"Riferimento normativo: {d.riferimento_normativo}")
                doc.add_paragraph("Fonti: " + "; ".join(f"{f.file} pag. {f.pagina}" for f in d.fonti))
                if req.note.get(d.id):
                    doc.add_paragraph(f"Nota del tecnico: {req.note[d.id]}")

        doc.add_paragraph()
        doc.add_paragraph(f"Il tecnico: {req.tecnico or '________________________'}")
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()


pratiche_service = PraticheService()
